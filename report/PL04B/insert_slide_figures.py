#!/usr/bin/env python3
"""Chèn hình từ report/figures vào deck bảo vệ SFDCT.
- Chèn vào dải trống các slide 9/13/15 (có chỗ sạch).
- Thêm slide hình chuyên dụng sau slide 7/8/10 cho các hình kiến trúc lớn.
Không ghi đè bản gốc: xuất ra ..._with_figures.pptx.
"""
import copy, struct, os
from pptx import Presentation
from pptx.util import Emu, Pt

SRC = "/home/huanthuytnhh/Downloads/SFDCT_ A Frequency-Aware Deepfake Detector for Banking eKYC.pptx"
OUT = "/home/huanthuytnhh/Desktop/thanhln/datn/report/SFDCT_defense_with_figures.pptx"
FIG = "/home/huanthuytnhh/Desktop/thanhln/datn/report/figures/"
EMU = 914400

def png_size(path):
    with open(path, "rb") as f:
        head = f.read(26)
    w, h = struct.unpack(">II", head[16:24])
    return w, h

def fit_box(path, box):
    """box=(l,t,r,b) inch. Trả (left,top,w,h) EMU giữ tỉ lệ, canh giữa."""
    l, t, r, b = [x * EMU for x in box]
    bw, bh = r - l, b - t
    iw, ih = png_size(path)
    ar = iw / ih
    w = bw; h = w / ar
    if h > bh:
        h = bh; w = h * ar
    left = l + (bw - w) / 2
    top = t + (bh - h) / 2
    return int(left), int(top), int(w), int(h)

def add_pic(slide, path, box):
    left, top, w, h = fit_box(path, box)
    slide.shapes.add_picture(path, left, top, w, h)

prs = Presentation(SRC)
W, H = prs.slide_width, prs.slide_height
def S(n): return prs.slides[n - 1]

# ---------- A) Chèn vào dải trống slide có sẵn chỗ ----------
add_pic(S(9),  FIG + "fig_3_13_gate_alpha.png", (9.35, 4.85, 17.35, 6.92))
add_pic(S(13), FIG + "fig_3_7_roc.png",          (2.6, 6.7, 8.8, 9.4))
add_pic(S(13), FIG + "fig_3_9_train_hff_r3.png", (8.9, 6.85, 15.8, 9.25))
add_pic(S(15), FIG + "fig_3_12_gradcam.png",     (4.9, 7.2, 13.0, 9.42))

# ---------- B) Slide hình chuyên dụng ----------
# Lấy "chrome" (thanh trang trí + tag chương + rule + title) từ slide 7 để khớp giao diện.
src7 = prs.slides[6]
chrome_src = {"leftbar": None, "topbar": None, "tag": None, "rule": None, "title": None}
for sh in src7.shapes:
    l = round(sh.left / EMU, 2); t = round(sh.top / EMU, 2)
    w = round(sh.width / EMU, 2); h = round(sh.height / EMU, 2)
    txt = sh.text_frame.text if sh.has_text_frame else ""
    if l == 0 and t == 0 and w < 0.2 and h > 9:        chrome_src["leftbar"] = sh._element
    elif l == 0 and t == 0 and w > 17 and h < 0.1:     chrome_src["topbar"] = sh._element
    elif "Methodology" in txt and l > 12:              chrome_src["tag"] = sh._element
    elif l < 0.6 and 0.9 < t < 1.0 and w > 16:         chrome_src["rule"] = sh._element
    elif "Frequency Hypothesis and Two" in txt:        chrome_src["title"] = sh._element

def set_title_text(title_el, text):
    """Đặt text cho element title đã deep-copy (giữ font/màu run đầu)."""
    ts = title_el.findall(".//" + "{http://schemas.openxmlformats.org/drawingml/2006/main}t")
    if ts:
        ts[0].text = text
        for extra in ts[1:]:
            extra.text = ""

def new_fig_slide(title, pics):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    spTree = slide.shapes._spTree
    for key in ("leftbar", "topbar", "tag", "rule", "title"):
        el = chrome_src[key]
        if el is None: continue
        cp = copy.deepcopy(el)
        if key == "title":
            set_title_text(cp, title)
        spTree.append(cp)
    for path, box in pics:
        add_pic(slide, path, box)
    return slide

# Tạo theo ĐÚNG thứ tự để map sldId: data, footprint, arch, zigzag, hff, liveness
new_fig_slide("Data Processing Pipeline", [
    (FIG + "fig_2_preprocess_pipeline.png",    (0.8, 1.45, 17.0, 4.9)),
    (FIG + "fig_3_2_preprocess_realfake.png",  (4.2, 4.85, 13.6, 9.4)),
])
new_fig_slide("Forgery Footprints — FF++ Manipulations", [
    (FIG + "fig_ff_manipulations.png", (0.7, 1.4, 17.1, 9.35)),
])
new_fig_slide("SFDCT — Two-Stream Architecture", [
    (FIG + "fig_sfdct_architecture_final.png", (0.6, 1.35, 17.2, 6.75)),
    (FIG + "fig_3_11_frequency.png",           (4.6, 6.55, 13.2, 9.42)),
])
new_fig_slide("Block-DCT Descriptor — Zigzag to 48-D", [
    (FIG + "fig_1_3_zigzag_detailed.png", (1.0, 1.5, 16.8, 9.3)),
])
new_fig_slide("SFDCT-HFF & Liveness Reuse", [
    (FIG + "fig_sfdct_hff_architect_final.png", (1.8, 1.4, 16.0, 5.2)),
    (FIG + "fig_2_b4dct_liveness.png",          (0.7, 5.45, 8.7, 9.3)),
    (FIG + "fig_hff_residual.png",              (9.1, 5.45, 17.1, 9.3)),
])
new_fig_slide("Liveness Detection (Secondary Module)", [
    (FIG + "fig_3_17_liveness.png",         (1.4, 1.4, 16.4, 5.55)),
    (FIG + "fig_3_20_roc_b4_liveness.png",  (2.4, 5.65, 15.4, 9.35)),
])

# ---------- C) Sắp lại thứ tự deck ----------
sldIdLst = prs.slides._sldIdLst
sldIds = list(sldIdLst)          # 17 gốc + 6 mới (cuối, theo thứ tự tạo)
orig = sldIds[:17]
dp, fp, arch, zz, hf, lv = sldIds[17:23]
desired = (
    orig[0:6]                    # 1 Title .. 6 Ch2 divider
    + [dp, fp]                   # NEW data pipeline, NEW footprint
    + [orig[6]]                  # Hypothesis + two-stream (gốc slide 7)
    + [arch]                     # NEW SFDCT architecture
    + [orig[7]]                  # Block-DCT descriptor (gốc 8)
    + [zz]                       # NEW zigzag
    + [orig[8]]                  # Gated attention (gốc 9)
    + [orig[9]]                  # SFDCT-HFF + liveness text (gốc 10)
    + [hf, lv]                   # NEW HFF figs, NEW liveness
    + orig[10:17]                # Ch3 div .. Thank you
)
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in desired:
    sldIdLst.append(el)

prs.save(OUT)
print("Saved:", OUT)
print("Total slides:", len(list(sldIdLst)))
print("chrome found:", {k: (v is not None) for k, v in chrome_src.items()})
